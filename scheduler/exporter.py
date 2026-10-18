"""Exportación a Excel (.xlsx) con openpyxl. Abre en Excel y en LibreOffice.

Seguridad: ningún texto proveniente del usuario se escribe como fórmula. Si empieza
con =, +, - o @ se guarda explícitamente como texto.
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import NOMBRE_DIA, Problema
from .service import Revision, Salida
from .vistas import cuadricula, entidades, etiqueta_materia, etiqueta_profesor

TEAL = "0D7377"
MINT = "A8D5BA"
CREMA = "F7F3E9"
OSCURO = "264653"

_FUENTE = "Calibri"
_BORDE = Side(style="thin", color="C9D6D3")
_BORDES = Border(left=_BORDE, right=_BORDE, top=_BORDE, bottom=_BORDE)
_PELIGRO = ("=", "+", "-", "@", "\t", "\r")


def _poner(ws, fila: int, col: int, valor, **estilo):
    """Escribe una celda tratando cualquier texto como texto (nunca como fórmula)."""
    c = ws.cell(row=fila, column=col)
    c.value = valor
    if isinstance(valor, str) and valor.startswith(_PELIGRO):
        c.data_type = "s"
    c.font = estilo.get("font", Font(name=_FUENTE, color=OSCURO))
    if "fill" in estilo:
        c.fill = estilo["fill"]
    if "alineacion" in estilo:
        c.alignment = estilo["alineacion"]
    if estilo.get("borde", False):
        c.border = _BORDES
    return c


def _encabezado(ws, fila: int, titulos: list[str], col0: int = 1):
    for i, t in enumerate(titulos):
        _poner(
            ws,
            fila,
            col0 + i,
            t,
            font=Font(name=_FUENTE, bold=True, color="FFFFFF"),
            fill=PatternFill("solid", fgColor=TEAL),
            alineacion=Alignment(horizontal="center", vertical="center", wrap_text=True),
            borde=True,
        )


def _anchos(ws, anchos: dict[int, float]):
    for col, ancho in anchos.items():
        ws.column_dimensions[get_column_letter(col)].width = ancho


def _hoja_general(wb: Workbook, salida: Salida):
    datos = salida.revision.datos
    ws = wb.active
    ws.title = "HORARIO_GENERAL"
    _encabezado(ws, 1, ["Día", "Hora_inicio", "Hora_fin", "Grupo", "Materia", "Profesor"])
    for i, a in enumerate(salida.asignaciones, start=2):
        valores = [
            NOMBRE_DIA.get(a.dia, a.dia),
            f"{a.inicio:02d}:00",
            f"{a.fin:02d}:00",
            a.grupo,
            etiqueta_materia(datos, a.materia),
            etiqueta_profesor(datos, a.profesor_id),
        ]
        for j, v in enumerate(valores, start=1):
            _poner(ws, i, j, v, borde=True)
    ws.freeze_panes = "A2"
    _anchos(ws, {1: 13, 2: 12, 3: 10, 4: 12, 5: 28, 6: 26})


def _hoja_cuadriculas(wb: Workbook, salida: Salida, tipo: str, nombre_hoja: str):
    datos = salida.revision.datos
    ws = wb.create_sheet(nombre_hoja)
    n_cols = len(datos.config.dias) + 1
    relleno = PatternFill("solid", fgColor=MINT)
    fila_actual = 1
    for etiqueta, valor in entidades(datos, salida.asignaciones, tipo):
        ws.merge_cells(start_row=fila_actual, start_column=1, end_row=fila_actual, end_column=n_cols)
        _poner(
            ws,
            fila_actual,
            1,
            etiqueta,
            font=Font(name=_FUENTE, bold=True, size=13, color=OSCURO),
            alineacion=Alignment(horizontal="left", vertical="center"),
        )
        encabezados, filas, ocupaciones = cuadricula(datos, salida.asignaciones, tipo, valor)
        _encabezado(ws, fila_actual + 1, encabezados)
        base = fila_actual + 2
        for i, fila in enumerate(filas):
            _poner(
                ws,
                base + i,
                1,
                fila[0],
                font=Font(name=_FUENTE, bold=True, color=OSCURO),
                alineacion=Alignment(horizontal="center", vertical="center"),
                borde=True,
            )
            for j in range(1, n_cols):
                _poner(ws, base + i, 1 + j, "", borde=True)
        for (i, j), (texto, dur) in ocupaciones.items():
            _poner(
                ws,
                base + i,
                1 + j,
                texto,
                fill=relleno,
                alineacion=Alignment(horizontal="center", vertical="center", wrap_text=True),
                borde=True,
            )
            if dur > 1:
                ws.merge_cells(start_row=base + i, start_column=1 + j, end_row=base + i + dur - 1, end_column=1 + j)
            for k in range(dur):
                ws.row_dimensions[base + i + k].height = 30
        fila_actual = base + len(filas) + 2
    _anchos(ws, {1: 14, **{c: 24 for c in range(2, n_cols + 1)}})


def _hoja_resumen(wb: Workbook, salida: Salida):
    datos = salida.revision.datos
    m = salida.metricas
    ws = wb.create_sheet("RESUMEN")
    ws.sheet_view.showGridLines = False
    estado_txt = "Horario óptimo" if salida.estado_motor == "OPTIMAL" else "Horario válido (no se demostró que sea el mejor posible)"
    filas = [
        ("Estado", estado_txt),
        ("Estado del optimizador", salida.estado_motor),
        ("Sesiones programadas", m.get("sesiones", 0)),
        ("Tiempo de búsqueda (s)", round(salida.segundos, 1)),
        ("Huecos internos de profesores (h)", m.get("huecos_profesores", 0)),
        ("Huecos internos de grupos (h)", m.get("huecos_grupos", 0)),
        ("Días de profesor sin huecos", f"{m.get('continuidad_pct', 100)} %"),
    ]
    r = 1
    for k, v in filas:
        _poner(ws, r, 1, k, font=Font(name=_FUENTE, bold=True, color=OSCURO))
        _poner(ws, r, 2, v)
        r += 1
    r += 1
    _poner(ws, r, 1, "Horas por profesor", font=Font(name=_FUENTE, bold=True, size=12, color=OSCURO))
    r += 1
    _encabezado(ws, r, ["Profesor", "Horas objetivo", "Horas asignadas", "Días con clase", "Huecos internos (h)"])
    for f in m.get("profesores", []):
        r += 1
        vals = [f["nombre"], f["horas_objetivo"] if f["horas_objetivo"] is not None else "—", f["horas"], f["dias"], f["huecos"]]
        for j, v in enumerate(vals, start=1):
            _poner(ws, r, j, v, borde=True)
    r += 2
    _poner(ws, r, 1, "Horas por grupo", font=Font(name=_FUENTE, bold=True, size=12, color=OSCURO))
    r += 1
    _encabezado(ws, r, ["Grupo", "Horas", "Huecos internos (h)"])
    for f in m.get("grupos", []):
        r += 1
        for j, v in enumerate([f["grupo"], f["horas"], f["huecos"]], start=1):
            _poner(ws, r, j, v, borde=True)
    _anchos(ws, {1: 36, 2: 22, 3: 18, 4: 16, 5: 20})


def _hoja_diagnostico(wb: Workbook, estado: str, problemas: list[Problema], avisos: list[str], sugerencias: list[str]):
    ws = wb.create_sheet("DIAGNOSTICO")
    _encabezado(ws, 1, ["Tipo", "Hoja", "Fila", "Detalle"])
    r = 2
    filas: list[tuple] = [("Estado", "", "", estado)]
    for p in problemas:
        filas.append(("Error" if p.es_error else "Advertencia", p.hoja, p.fila if p.fila is not None else "", p.mensaje))
    for a in avisos:
        filas.append(("Aviso", "", "", a))
    for s in sugerencias:
        filas.append(("Recomendación", "", "", s))
    for fila in filas:
        for j, v in enumerate(fila, start=1):
            _poner(ws, r, j, v, borde=True, alineacion=Alignment(vertical="top", wrap_text=(j == 4)))
        r += 1
    ws.freeze_panes = "A2"
    _anchos(ws, {1: 16, 2: 20, 3: 8, 4: 110})


def _guardar(wb: Workbook, ruta: str):
    wb.properties.creator = "Schedora"
    wb.properties.title = "Horario generado con Schedora"
    wb.save(ruta)


def exportar_horario(salida: Salida, ruta: str) -> str:
    """Exporta el horario. Solo si el resultado fue verificado (estado LISTO)."""
    if not salida.exitoso or salida.revision is None:
        raise ValueError("Solo se puede exportar un horario verificado.")
    wb = Workbook()
    _hoja_general(wb, salida)
    _hoja_cuadriculas(wb, salida, "grupo", "HORARIO_POR_GRUPO")
    _hoja_cuadriculas(wb, salida, "profesor", "HORARIO_POR_PROFESOR")
    _hoja_resumen(wb, salida)
    problemas = salida.revision.advertencias
    if problemas or salida.avisos or salida.sugerencias:
        _hoja_diagnostico(wb, f"{salida.titulo} ({salida.estado_motor})", problemas, salida.avisos, salida.sugerencias)
    _guardar(wb, ruta)
    return ruta


def exportar_diagnostico(rev: Revision | None, salida: Salida | None, ruta: str) -> str:
    """Reporte de diagnóstico: errores de entrada, advertencias, estado y recomendaciones."""
    wb = Workbook()
    wb.remove(wb.active)
    if salida is not None:
        estado = f"{salida.titulo} ({salida.estado_motor})"
        rev = salida.revision or rev
        avisos, sugerencias = salida.avisos, salida.sugerencias
    else:
        estado = "Revisión del archivo"
        avisos, sugerencias = [], []
    problemas = rev.problemas if rev is not None else []
    _hoja_diagnostico(wb, estado, problemas, avisos, sugerencias)
    _guardar(wb, ruta)
    return ruta
