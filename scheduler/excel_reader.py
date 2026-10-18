"""Lectura del libro Excel de entrada.

- Solo acepta .xlsx (nunca .xlsm: no se ejecutan macros).
- Lee valores, no fórmulas (data_only=True).
- Limita tamaño del archivo, tamaño descomprimido y número de filas.
- Devuelve los datos junto con una lista de Problema con hoja y fila.
"""
from __future__ import annotations

import datetime as dt
import os
import re
import unicodedata
import zipfile

from openpyxl import load_workbook

from .models import (
    DIAS,
    Config,
    Datos,
    ErrorArchivo,
    FilaDisponibilidad,
    Materia,
    MISMA_MATERIA_MODOS,
    Oferta,
    Problema,
    Profesor,
)

MAX_BYTES_ARCHIVO = 5 * 1024 * 1024
MAX_BYTES_DESCOMPRIMIDO = 60 * 1024 * 1024
MAX_FILAS_POR_HOJA = 5000

REQUERIDAS = {
    "PROFESORES": ["profesor_id", "nombre", "materias"],
    "MATERIAS": ["materia_id", "nombre_materia", "duracion_sesion"],
    "OFERTA_ACADEMICA": ["oferta_id", "grupo_id", "materia_id", "horas_semanales"],
}
OPCIONALES = ("DISPONIBILIDAD", "CONFIGURACION")


# --------------------------------------------------------------------------
# utilidades de normalización
# --------------------------------------------------------------------------
def _sin_acentos(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    return "".join(c for c in s if not unicodedata.combining(c))


def norm(s) -> str:
    """Normaliza nombres de hojas, columnas y días: minúsculas, sin acentos, con guiones bajos."""
    if s is None:
        return ""
    return re.sub(r"\s+", "_", _sin_acentos(str(s)).strip().lower())


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _entero(v, nombre: str) -> int:
    if isinstance(v, bool):
        raise ValueError(f"{nombre} debe ser un número entero")
    if isinstance(v, (int, float)):
        if float(v).is_integer():
            return int(v)
        raise ValueError(f"{nombre} debe ser un número entero (se recibió {v})")
    t = _texto(v)
    if re.fullmatch(r"-?\d+", t):
        return int(t)
    if re.fullmatch(r"-?\d+[.,]0+", t):
        return int(re.split(r"[.,]", t)[0])
    raise ValueError(f"{nombre} debe ser un número entero (se recibió «{t}»)")


def _numero(v, nombre: str) -> float:
    if isinstance(v, bool):
        raise ValueError(f"{nombre} debe ser un número")
    if isinstance(v, (int, float)):
        return float(v)
    t = _texto(v).replace(",", ".")
    try:
        return float(t)
    except ValueError:
        raise ValueError(f"{nombre} debe ser un número (se recibió «{_texto(v)}»)") from None


def _hora(v, nombre: str) -> int:
    """Convierte una celda de hora a entero. Solo horas en punto."""
    if isinstance(v, dt.datetime):
        v = v.time()
    if isinstance(v, dt.time):
        h, m = v.hour, v.minute
    elif isinstance(v, bool):
        raise ValueError(f"{nombre} no es una hora válida")
    elif isinstance(v, (int, float)):
        x = float(v)
        if 0 < x < 1:  # fracción de día de Excel
            minutos = round(x * 24 * 60)
            h, m = divmod(minutos, 60)
        elif x.is_integer() and 0 <= x <= 24:
            h, m = int(x), 0
        else:
            raise ValueError(f"{nombre} no es una hora válida (se recibió {v})")
    else:
        t = _texto(v)
        mt = re.fullmatch(r"(\d{1,2})(?::(\d{2}))?(?::\d{2})?", t)
        if not mt:
            raise ValueError(f"{nombre} no es una hora válida (se recibió «{t}»); use el formato 07:00")
        h, m = int(mt.group(1)), int(mt.group(2) or 0)
    if m != 0:
        raise ValueError(f"{nombre} debe ser una hora en punto (07:00, 08:00…); se recibió {h:02d}:{m:02d}")
    if not 0 <= h <= 24:
        raise ValueError(f"{nombre} fuera de rango (se recibió {h}:00)")
    return h


def _lista(v) -> list[str]:
    t = _texto(v)
    if not t:
        return []
    return [p.strip() for p in re.split(r"[;,\n]", t) if p.strip()]


def _dia(v, nombre: str = "dia") -> str:
    d = norm(v)
    if d not in DIAS:
        raise ValueError(f"{nombre}: «{_texto(v)}» no es un día válido (lunes a sábado)")
    return d


def _lista_dias(v) -> list[str]:
    t = norm(v).replace("_a_", " a ")
    # rango "lunes a viernes"
    mt = re.fullmatch(r"([a-z]+) a ([a-z]+)", t.replace("_", " "))
    if mt and mt.group(1) in DIAS and mt.group(2) in DIAS:
        i, j = DIAS.index(mt.group(1)), DIAS.index(mt.group(2))
        if i <= j:
            return DIAS[i : j + 1]
    return [_dia(p) for p in _lista(v)]


def _booleano(v, defecto: bool = True) -> bool:
    t = norm(v)
    if t == "":
        return defecto
    if t in ("no", "false", "0", "inactivo", "n", "falso"):
        return False
    if t in ("si", "sí", "true", "1", "activo", "s", "verdadero"):
        return True
    raise ValueError(f"«{_texto(v)}» no es un valor sí/no válido")


# --------------------------------------------------------------------------
# apertura segura del libro
# --------------------------------------------------------------------------
def abrir_libro(ruta: str):
    if not ruta or not os.path.exists(ruta):
        raise ErrorArchivo("No se encontró el archivo. Vuelve a seleccionarlo e inténtalo de nuevo.")
    nombre = os.path.basename(ruta).lower()
    if nombre.endswith(".xlsm") or nombre.endswith(".xlsb") or nombre.endswith(".xls"):
        raise ErrorArchivo(
            "Solo se aceptan archivos .xlsx. Guarda tu libro como «Libro de Excel (.xlsx)» y vuelve a subirlo."
        )
    if not nombre.endswith(".xlsx"):
        raise ErrorArchivo("El archivo debe tener extensión .xlsx.")
    if os.path.getsize(ruta) > MAX_BYTES_ARCHIVO:
        raise ErrorArchivo(
            f"El archivo es demasiado grande (máximo {MAX_BYTES_ARCHIVO // (1024 * 1024)} MB)."
        )
    try:
        with zipfile.ZipFile(ruta) as z:
            total = sum(i.file_size for i in z.infolist())
            if total > MAX_BYTES_DESCOMPRIMIDO:
                raise ErrorArchivo("El archivo contiene demasiada información y no se puede procesar.")
            if "xl/workbook.xml" not in z.namelist():
                raise ErrorArchivo("El archivo no parece ser un libro de Excel válido.")
    except zipfile.BadZipFile:
        raise ErrorArchivo("El archivo no parece ser un libro de Excel válido (.xlsx).") from None
    try:
        return load_workbook(ruta, read_only=True, data_only=True, keep_links=False)
    except ErrorArchivo:
        raise
    except Exception:
        raise ErrorArchivo(
            "No se pudo abrir el libro. Verifica que no esté dañado ni protegido con contraseña."
        ) from None


def _leer_hoja(ws, hoja: str, requeridas: list[str], problemas: list[Problema]):
    """Devuelve [(num_fila_excel, {columna_normalizada: valor})]."""
    filas = ws.iter_rows(values_only=True)
    idx = 0
    encabezado = None
    for fila in filas:
        idx += 1
        if any(c is not None and str(c).strip() != "" for c in fila):
            encabezado = [norm(c) for c in fila]
            break
    if encabezado is None:
        if requeridas:
            problemas.append(Problema("error", hoja, None, "La hoja está vacía."))
        return []
    faltantes = [c for c in requeridas if c not in encabezado]
    if faltantes:
        problemas.append(
            Problema("error", hoja, idx, "Faltan columnas obligatorias: " + ", ".join(faltantes) + ".")
        )
        return []
    registros = []
    for fila in filas:
        idx += 1
        if all(c is None or str(c).strip() == "" for c in fila):
            continue
        if len(registros) >= MAX_FILAS_POR_HOJA:
            problemas.append(
                Problema("error", hoja, idx, f"La hoja supera el máximo de {MAX_FILAS_POR_HOJA} filas.")
            )
            break
        registros.append((idx, dict(zip(encabezado, fila))))
    return registros


def _buscar_hojas(wb) -> dict[str, object]:
    return {norm(n).upper(): wb[n] for n in wb.sheetnames}


# --------------------------------------------------------------------------
# lectura principal
# --------------------------------------------------------------------------
def leer_excel(ruta: str) -> tuple[Datos, list[Problema]]:
    """Lee el libro y devuelve (Datos, problemas). Lanza ErrorArchivo si no se puede abrir."""
    wb = abrir_libro(ruta)
    problemas: list[Problema] = []
    datos = Datos()
    try:
        hojas = _buscar_hojas(wb)
        for nombre in REQUERIDAS:
            if nombre not in hojas:
                problemas.append(Problema("error", nombre, None, f"Falta la hoja «{nombre}»."))
        if any(p.es_error for p in problemas):
            return datos, problemas

        _config(hojas.get("CONFIGURACION"), datos, problemas)
        _profesores(hojas["PROFESORES"], datos, problemas)
        _materias(hojas["MATERIAS"], datos, problemas)
        _ofertas(hojas["OFERTA_ACADEMICA"], datos, problemas)
        if "DISPONIBILIDAD" in hojas:
            _disponibilidad(hojas["DISPONIBILIDAD"], datos, problemas)
    finally:
        wb.close()
    return datos, problemas


def _config(ws, datos: Datos, problemas: list[Problema]) -> None:
    if ws is None:
        return
    cfg = datos.config
    for fila, reg in _leer_hoja(ws, "CONFIGURACION", ["clave", "valor"], problemas):
        clave, valor = norm(reg.get("clave")), reg.get("valor")
        if clave == "":
            continue
        try:
            if clave == "hora_inicio":
                cfg.hora_inicio = _hora(valor, "hora_inicio")
            elif clave == "hora_fin":
                cfg.hora_fin = _hora(valor, "hora_fin")
            elif clave == "dias":
                cfg.dias = _lista_dias(valor)
            elif clave == "limite_busqueda_segundos":
                cfg.limite_busqueda_segundos = _entero(valor, clave)
            elif clave == "peso_continuidad":
                cfg.peso_continuidad = _entero(valor, clave)
            elif clave in ("peso_grupos", "peso_fragmentacion"):
                cfg.peso_grupos = _entero(valor, clave)
            elif clave == "peso_dias_docente":
                cfg.peso_dias_docente = _entero(valor, clave)
            elif clave in ("peso_misma_materia_dia", "peso_distribucion"):
                cfg.peso_misma_materia_dia = _entero(valor, clave)
            elif clave == "misma_materia_mismo_dia":
                m = norm(valor)
                if m not in MISMA_MATERIA_MODOS:
                    raise ValueError("debe ser: " + ", ".join(MISMA_MATERIA_MODOS))
                cfg.misma_materia_mismo_dia = m
            elif clave == "semilla_busqueda":
                cfg.semilla_busqueda = None if _texto(valor) == "" else _entero(valor, clave)
            else:
                problemas.append(
                    Problema("advertencia", "CONFIGURACION", fila, f"Parámetro desconocido «{clave}»; se ignora.")
                )
        except ValueError as e:
            problemas.append(Problema("error", "CONFIGURACION", fila, str(e)))
    if not cfg.dias:
        problemas.append(Problema("error", "CONFIGURACION", None, "La lista de días no puede estar vacía."))
    if cfg.hora_inicio >= cfg.hora_fin:
        problemas.append(
            Problema("error", "CONFIGURACION", None, "hora_inicio debe ser anterior a hora_fin.")
        )


def _profesores(ws, datos: Datos, problemas: list[Problema]) -> None:
    for fila, r in _leer_hoja(ws, "PROFESORES", REQUERIDAS["PROFESORES"], problemas):
        try:
            pid = _texto(r.get("profesor_id"))
            if not pid:
                raise ValueError("profesor_id está vacío")
            p = Profesor(
                id=pid,
                nombre=_texto(r.get("nombre")) or pid,
                materias=set(_lista(r.get("materias"))),
                activo=_booleano(r.get("activo")),
            )
            if _texto(r.get("horas_objetivo_semana")):
                p.horas_objetivo = _numero(r.get("horas_objetivo_semana"), "horas_objetivo_semana")
            if _texto(r.get("horas_maximas_semana")):
                p.horas_maximas = _entero(r.get("horas_maximas_semana"), "horas_maximas_semana")
            if _texto(r.get("horas_minimas_semana")):
                p.horas_minimas = _entero(r.get("horas_minimas_semana"), "horas_minimas_semana")
            datos.profesores.append(p)
        except ValueError as e:
            problemas.append(Problema("error", "PROFESORES", fila, str(e)))


def _materias(ws, datos: Datos, problemas: list[Problema]) -> None:
    for fila, r in _leer_hoja(ws, "MATERIAS", REQUERIDAS["MATERIAS"], problemas):
        try:
            mid = _texto(r.get("materia_id"))
            if not mid:
                raise ValueError("materia_id está vacío")
            datos.materias.append(
                Materia(
                    id=mid,
                    nombre=_texto(r.get("nombre_materia")) or mid,
                    duracion=_entero(r.get("duracion_sesion"), "duracion_sesion"),
                )
            )
        except ValueError as e:
            problemas.append(Problema("error", "MATERIAS", fila, str(e)))


def _ofertas(ws, datos: Datos, problemas: list[Problema]) -> None:
    for fila, r in _leer_hoja(ws, "OFERTA_ACADEMICA", REQUERIDAS["OFERTA_ACADEMICA"], problemas):
        try:
            oid = _texto(r.get("oferta_id"))
            grupo = _texto(r.get("grupo_id"))
            if not oid:
                raise ValueError("oferta_id está vacío")
            if not grupo:
                raise ValueError("grupo_id está vacío")
            o = Oferta(
                id=oid,
                grupo=grupo,
                materia=_texto(r.get("materia_id")),
                horas_semanales=_entero(r.get("horas_semanales"), "horas_semanales"),
            )
            if _texto(r.get("sesiones_semanales")):
                o.sesiones_semanales = _entero(r.get("sesiones_semanales"), "sesiones_semanales")
            if _texto(r.get("profesor_obligatorio")):
                o.profesor_obligatorio = _texto(r.get("profesor_obligatorio"))
            if _texto(r.get("dias_permitidos")):
                o.dias_permitidos = _lista_dias(r.get("dias_permitidos"))
            datos.ofertas.append(o)
        except ValueError as e:
            problemas.append(Problema("error", "OFERTA_ACADEMICA", fila, str(e)))


def _disponibilidad(ws, datos: Datos, problemas: list[Problema]) -> None:
    req = ["profesor_id", "dia", "hora_inicio", "hora_fin", "tipo"]
    for fila, r in _leer_hoja(ws, "DISPONIBILIDAD", req, problemas):
        try:
            pid = _texto(r.get("profesor_id"))
            if not pid:
                raise ValueError("profesor_id está vacío")
            tipo = norm(r.get("tipo")).upper()
            if tipo not in ("DISPONIBLE", "NO_DISPONIBLE"):
                raise ValueError(f"tipo «{_texto(r.get('tipo'))}» no es válido; use DISPONIBLE o NO_DISPONIBLE")
            datos.disponibilidad.append(
                FilaDisponibilidad(
                    profesor_id=pid,
                    dia=_dia(r.get("dia")),
                    hora_inicio=_hora(r.get("hora_inicio"), "hora_inicio"),
                    hora_fin=_hora(r.get("hora_fin"), "hora_fin"),
                    tipo=tipo,
                    fila=fila,
                )
            )
        except ValueError as e:
            problemas.append(Problema("error", "DISPONIBILIDAD", fila, str(e)))
