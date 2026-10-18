import pytest
from openpyxl import load_workbook

from helpers import ejemplo, horario_voraz
from scheduler.exporter import exportar_diagnostico, exportar_horario
from scheduler.metrics import calcular_metricas
from scheduler.service import Salida, revisar


@pytest.fixture(scope="module")
def salida_lista():
    rev = revisar(ejemplo("datos_validos.xlsx"))
    a = horario_voraz(rev.datos, rev.sesiones, rev.disp)
    return Salida("LISTO", "FEASIBLE", "Tu horario está listo", "ok", revision=rev, asignaciones=a,
                  metricas=calcular_metricas(rev.datos, a), segundos=1.5)


def test_hojas_del_libro_exportado(tmp_path, salida_lista):
    ruta = exportar_horario(salida_lista, str(tmp_path / "h.xlsx"))
    wb = load_workbook(ruta)
    assert wb.sheetnames[:4] == ["HORARIO_GENERAL", "HORARIO_POR_GRUPO", "HORARIO_POR_PROFESOR", "RESUMEN"]


def test_horario_general_coincide_con_la_solucion_validada(tmp_path, salida_lista):
    ruta = exportar_horario(salida_lista, str(tmp_path / "h.xlsx"))
    ws = load_workbook(ruta)["HORARIO_GENERAL"]
    filas = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(filas) == len(salida_lista.asignaciones) == 48
    assert [c.value for c in ws[1]] == ["Día", "Hora_inicio", "Hora_fin", "Grupo", "Materia", "Profesor"]
    horas = sum(int(f[2][:2]) - int(f[1][:2]) for f in filas)
    assert horas == 76


def test_cuadricula_por_grupo_combina_celdas_de_dos_horas(tmp_path, salida_lista):
    ruta = exportar_horario(salida_lista, str(tmp_path / "h.xlsx"))
    ws = load_workbook(ruta)["HORARIO_POR_GRUPO"]
    assert ws.merged_cells.ranges  # las sesiones de 2 h y los títulos están combinados
    textos = [c.value for fila in ws.iter_rows() for c in fila if isinstance(c.value, str)]
    assert any("Matemáticas · Ana Torres" in t for t in textos)


def test_resumen_con_estado_y_horas(tmp_path, salida_lista):
    ruta = exportar_horario(salida_lista, str(tmp_path / "h.xlsx"))
    ws = load_workbook(ruta)["RESUMEN"]
    valores = [c.value for fila in ws.iter_rows() for c in fila if c.value is not None]
    assert "Sesiones programadas" in valores and 48 in valores and "Ana Torres" in valores


def test_no_se_exporta_un_horario_no_verificado(tmp_path, salida_lista):
    malo = Salida("NO_VERIFICADO", "FEASIBLE", "x", "y", revision=salida_lista.revision, asignaciones=salida_lista.asignaciones)
    with pytest.raises(ValueError):
        exportar_horario(malo, str(tmp_path / "h.xlsx"))
    assert not (tmp_path / "h.xlsx").exists()


def test_texto_que_parece_formula_se_guarda_como_texto(tmp_path, salida_lista):
    rev = salida_lista.revision
    original = rev.datos.profesor("P01").nombre
    rev.datos.profesor("P01").nombre = '=HYPERLINK("http://malo.example","clic")'
    try:
        ruta = exportar_horario(salida_lista, str(tmp_path / "h.xlsx"))
    finally:
        rev.datos.profesor("P01").nombre = original
    ws = load_workbook(ruta)["HORARIO_GENERAL"]
    celdas = [c for fila in ws.iter_rows(min_row=2) for c in fila if isinstance(c.value, str) and c.value.startswith("=")]
    assert celdas and all(c.data_type == "s" for c in celdas)  # texto, nunca fórmula ("f")


def test_diagnostico_de_entrada_invalida(tmp_path):
    rev = revisar(ejemplo("datos_infactibles.xlsx"))
    ruta = exportar_diagnostico(rev, None, str(tmp_path / "d.xlsx"))
    wb = load_workbook(ruta)
    assert wb.sheetnames == ["DIAGNOSTICO"]
    texto = " ".join(str(c.value) for fila in wb["DIAGNOSTICO"].iter_rows() for c in fila if c.value)
    assert "Educación Física" in texto and "Error" in texto


def test_diagnostico_incluye_recomendaciones(tmp_path):
    rev = revisar(ejemplo("datos_validos.xlsx"))
    s = Salida("SIN_SOLUCION", "INFEASIBLE", "No existe un horario con estas reglas", "m", revision=rev,
               sugerencias=["Si se relaja X sí existe"])
    wb = load_workbook(exportar_diagnostico(rev, s, str(tmp_path / "d.xlsx")))
    texto = " ".join(str(c.value) for fila in wb["DIAGNOSTICO"].iter_rows() for c in fila if c.value)
    assert "Recomendación" in texto and "INFEASIBLE" in texto
